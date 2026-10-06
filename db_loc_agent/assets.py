"""结构化迁移资产服务：资产落库、报告模板渲染、Excel 导出。

区别于纯文本 markdown 产物：
- 每个风险项结构化存储（对象名称/风险等级/不兼容点/原始SQL/转换SQL/改造建议/来源/状态/负责人）；
- 最终报告由后端读取资产表按模板渲染生成，LLM 只负责填充业务描述；
- 支持导出 Excel 资产清单，供工程实施人员直接使用。
"""
from __future__ import annotations

import io
import json
from typing import Any, Dict, List, Optional

from db_loc_agent.db import DbMigrateAsset, DbMigrateTask

_ASSET_TYPES = ("table", "view", "procedure", "function", "trigger", "index", "sequence", "synonym")

# 报告模板（资产自动渲染主体 + LLM 填充业务描述的占位区）
REPORT_TEMPLATE = """# {task_name}——数据库国产化迁移评估报告

## 一、任务概要
- **任务状态**：{status_label}
- **源数据库**：{source_db_type}（数据源：{source_ds_name}）
- **目标数据库**：{target_db_type}
- **业务等级**：{db_business_level}｜**库容量**：{db_capacity}｜**停机窗口**：{stop_window}
- **扫描对象数**：{object_count}

## 二、兼容性评估结论（LLM 业务描述区）
{summary}

## 三、资产风险清单（结构化资产自动渲染，可导出 Excel）
| # | 对象 | 类型 | 风险 | 来源 | 状态 | 命中规则 |
|---|------|------|------|------|------|----------|
{asset_rows}

## 四、高风险对象明细与改造建议
{detail_rows}

## 五、转换 SQL 脚本
{convert_rows}
"""

_RISK_ORDER = {"高": 0, "中": 1, "低": 2}


class MigrateAssetService:
    """db_migrate_asset 表的服务：增删改查 + 报告渲染 + Excel 导出。"""

    def __init__(self, db):
        self.db = db
        self.rule_engine = None  # 在 app.py 装配时注入

    # ---------- 写入 ----------

    def upsert(self, task_db_id: int, asset: Dict[str, Any],
               conv_id: Optional[int] = None) -> Dict[str, Any]:
        """新增或更新资产（按 task_id + asset_type + asset_name 匹配更新）。"""
        from sqlalchemy import select
        asset_name = str(asset.get("asset_name") or "").strip()
        asset_type = str(asset.get("asset_type") or "table").strip()
        if not asset_name:
            raise ValueError("资产对象名称不能为空")
        if asset_type not in _ASSET_TYPES:
            asset_type = "table"
        with self.db.session() as s:
            row = s.scalar(select(DbMigrateAsset).where(
                DbMigrateAsset.task_id == task_db_id,
                DbMigrateAsset.asset_name == asset_name,
                DbMigrateAsset.asset_type == asset_type))
            if row is None:
                row = DbMigrateAsset(task_id=task_db_id, asset_type=asset_type,
                                     asset_name=asset_name, conv_id=conv_id)
                s.add(row)
            if asset.get("risk_level"):
                row.risk_level = str(asset["risk_level"])
            incompatible = asset.get("incompatible_detail")
            if incompatible is not None:
                row.incompatible_detail = incompatible if isinstance(incompatible, str) \
                    else json.dumps(incompatible, ensure_ascii=False)
            if asset.get("rule_ids") is not None:
                row.rule_ids = (asset["rule_ids"] if isinstance(asset["rule_ids"], str)
                                else json.dumps(asset["rule_ids"], ensure_ascii=False))
            if asset.get("original_sql") is not None:
                row.original_sql = asset["original_sql"]
            if asset.get("converted_sql") is not None:
                row.converted_sql = asset["converted_sql"]
            if asset.get("suggestion") is not None:
                row.suggestion = asset["suggestion"]
            if asset.get("source") is not None:
                row.source = asset["source"]
            if asset.get("status") is not None:
                row.status = asset["status"]
            if asset.get("owner") is not None:
                row.owner = asset["owner"]
            s.commit()
            s.refresh(row)
            return self._to_dict(row)

    def bulk_add(self, task_db_id: int, assets: List[Dict[str, Any]],
                 conv_id: Optional[int] = None) -> List[Dict[str, Any]]:
        """批量入库（规则引擎扫描结果直接写资产表），返回新增/更新后的资产列表。"""
        out = []
        for a in assets:
            try:
                out.append(self.upsert(task_db_id, a, conv_id=conv_id))
            except (ValueError, Exception):
                continue  # 单条失败跳过，不影响批量
        return out

    # ---------- 读取 ----------

    def list(self, task_db_id: int) -> List[Dict[str, Any]]:
        from sqlalchemy import select
        with self.db.session() as s:
            rows = list(s.scalars(select(DbMigrateAsset)
                                  .where(DbMigrateAsset.task_id == task_db_id)
                                  .order_by(DbMigrateAsset.id)))
        assets = [self._to_dict(r) for r in rows]
        assets.sort(key=lambda a: _RISK_ORDER.get(a.get("risk_level", "低"), 9))
        return assets

    def get(self, asset_db_id: int) -> Optional[Dict[str, Any]]:
        with self.db.session() as s:
            r = s.get(DbMigrateAsset, asset_db_id)
            return self._to_dict(r) if r else None

    def delete(self, asset_db_id: int) -> None:
        with self.db.session() as s:
            r = s.get(DbMigrateAsset, asset_db_id)
            if r:
                s.delete(r)
                s.commit()

    def delete_by_task(self, task_db_id: int) -> int:
        """按任务批量删除资产（重置任务时清空已生成资产）。返回删除条数。"""
        from sqlalchemy import select, delete
        with self.db.session() as s:
            rows = list(s.scalars(select(DbMigrateAsset).where(DbMigrateAsset.task_id == task_db_id)))
            for r in rows:
                s.delete(r)
            s.commit()
            return len(rows)

    @staticmethod
    def _to_dict(r: DbMigrateAsset) -> Dict[str, Any]:
        return {
            "id": r.id, "task_id": r.task_id, "asset_type": r.asset_type,
            "asset_name": r.asset_name, "risk_level": r.risk_level,
            "incompatible_detail": r.incompatible_detail,
            "rule_ids": r.rule_ids, "original_sql": r.original_sql,
            "converted_sql": r.converted_sql, "suggestion": r.suggestion,
            "source": r.source, "status": r.status, "owner": r.owner,
            "conv_id": r.conv_id,
            "created_at": r.created_at.isoformat() if r.created_at else None,
            "updated_at": r.updated_at.isoformat() if r.updated_at else None,
        }

    # ---------- 报告渲染 ----------

    def render_report(self, task: Dict[str, Any],
                      llm_summary: str = "") -> str:
        """由结构化资产生成 Markdown 报告（模板渲染，LLM 只负责填充业务描述 summary）。"""
        assets = self.list(int(task["task_id"]))
        if assets:
            asset_rows = "\n".join(
                f"| {i + 1} | {a['asset_name']} | {a['asset_type']} | {a['risk_level']} "
                f"| {a['source']} | {a['status']} | {self._rule_ids_text(a)} |"
                for i, a in enumerate(assets))
            detail_rows = "\n".join(
                f"- **{a['asset_type']} `{a['asset_name']}`（{a['risk_level']}风险）**：{a['suggestion'] or '待评估'}"
                for a in assets if a["risk_level"] in ("高", "中"))
            convert_rows = "\n".join(
                f"### {a['asset_type']} `{a['asset_name']}`\n```sql\n{a['converted_sql'] or a['original_sql'] or '-- 暂无转换脚本'}\n```"
                for a in assets if a.get("converted_sql"))
        else:
            asset_rows = "（暂无结构化资产，请先执行对象扫描与规则评估）"
            detail_rows = "（暂无高风险对象明细）"
            convert_rows = "（暂无转换 SQL 脚本）"
        return REPORT_TEMPLATE.format(
            task_name=task.get("task_name") or "数据库迁移评估",
            status_label=task.get("status_label") or task.get("status") or "",
            source_db_type=task.get("source_db_type") or "待定",
            source_ds_name=task.get("source_ds_name") or "",
            target_db_type=task.get("target_db_type") or "待定",
            db_business_level=task.get("db_business_level") or "待定",
            db_capacity=task.get("db_capacity") or "待定",
            stop_window=task.get("stop_window") or "待定",
            object_count=task.get("object_count") or 0,
            summary=llm_summary.strip() or "（由评估专家 Agent 补充业务影响解读，此处自动填充。）",
            asset_rows=asset_rows, detail_rows=detail_rows, convert_rows=convert_rows,
        )

    @staticmethod
    def _rule_ids_text(a: Dict[str, Any]) -> str:
        try:
            rid = json.loads(a["rule_ids"])
            if isinstance(rid, list):
                return "、".join(str(x) for x in rid)
        except Exception:
            pass
        return a.get("rule_ids") or "-"

    # ---------- Excel 导出 ----------

    def export_excel(self, task_db_id: int) -> bytes:
        """导出资产清单 Excel（openpyxl），供工程实施人员直接使用。"""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        assets = self.list(task_db_id)
        with self.db.session() as s:
            t = s.get(DbMigrateTask, task_db_id)
            task_name = t.task_name if t else "数据库迁移评估"
        wb = Workbook()
        ws = wb.active
        ws.title = "迁移资产清单"
        ws["A1"] = f"{task_name}——迁移资产清单"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:K1")
        headers = ["序号", "对象类型", "对象名称", "风险等级", "不兼容点",
                   "命中规则", "原始SQL/DDL", "转换SQL", "改造建议", "来源", "状态"]
        ws.append([])
        ws.append(headers)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="5B7CFA")
            c.alignment = Alignment(horizontal="center", vertical="center")
        for i, a in enumerate(assets, 1):
            try:
                detail = json.loads(a["incompatible_detail"])
                detail_text = "\n".join(
                    f"{d.get('rule_id', '')} {d.get('desc', '')}" for d in detail) \
                    if isinstance(detail, list) else str(detail)
            except Exception:
                detail_text = a.get("incompatible_detail") or ""
            ws.append([i, a["asset_type"], a["asset_name"], a["risk_level"],
                       detail_text, self._rule_ids_text(a),
                       a["original_sql"], a["converted_sql"], a["suggestion"],
                       a["source"], a["status"]])
        widths = [6, 10, 18, 8, 36, 12, 40, 40, 30, 12, 10]
        for idx, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + idx)].width = w
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ---------- 统计 ----------

    def stats(self, task_db_id: int) -> Dict[str, Any]:
        assets = self.list(task_db_id)
        return {
            "total": len(assets),
            "high": sum(1 for a in assets if a["risk_level"] == "高"),
            "medium": sum(1 for a in assets if a["risk_level"] == "中"),
            "low": sum(1 for a in assets if a["risk_level"] == "低"),
            "by_type": {t: sum(1 for a in assets if a["asset_type"] == t)
                        for t in _ASSET_TYPES},
            "by_source": {src: sum(1 for a in assets if a["source"] == src)
                          for src in ("rule_engine", "assessor", "transformer")},
        }