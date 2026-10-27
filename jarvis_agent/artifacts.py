"""产物写入与 namespace 管理"""
from __future__ import annotations

import re
import difflib
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Optional, Union

from sqlalchemy import select

from jarvis_agent.db import ArtifactRecord, Database

# namespace 各段允许：字母数字、下划线、连字符、中文（中文此前被拒绝，导致
# "浙江营业库A" 这类对象名进不了 namespace，safe_object_name 又把中文删成只剩 "A"）
SEGMENT_RE = re.compile(r"^[A-Za-z0-9_\u4e00-\u9fff-]{1,128}$")


class NamespaceError(ValueError):
    pass


def make_namespace(conv_key: str, agent_id: str, object_name: str,
                   stamp: Optional[str] = None) -> str:
    for seg in (conv_key, agent_id, object_name):
        if not SEGMENT_RE.match(seg or ""):
            raise NamespaceError(f"非法 namespace 段: {seg!r}")
    stamp = stamp or datetime.now().strftime("%Y%m%dT%H%M%S")
    if not SEGMENT_RE.match(stamp):
        raise NamespaceError("非法时间戳")
    return f"{conv_key}/{agent_id}/{object_name}/{stamp}"


def safe_object_name(text: str, max_len: int = 128) -> str:
    # 保留中文/字母数字/下划线/连字符，其余替换为下划线；此前 [^A-Za-z0-9_-] 会把
    # "浙江营业库A" 里的中文全删成只剩 "A"，导致对象名/目录名退化成单字符垃圾名
    cleaned = re.sub(r"[^A-Za-z0-9_\u4e00-\u9fff-]+", "_", (text or "").strip()).strip("_")
    return (cleaned or "artifact")[:max_len]


# 提取失败/无意义时回退的默认对象名；以及应识别为"非对象名"的垃圾值（参数键名、占位词）
_DEFAULT_OBJECT_NAME = "artifact"
_JUNK_OBJECT_NAMES = {
    "artifact", "value", "confirm", "task_confirm",
    "source_data_id", "target_data_id", "source_db", "target_db",
    "source_db_type", "target_db_type", "datasource", "object_name",
}


def primary_object_name(text: str, max_len: int = 40) -> str:
    """从需求/对象清单中取主对象名（用于 namespace）：多对象取第一项、过长的截断。

    - text 为空 -> 默认名
    - 逗号/空格/换行分隔的清单 -> 取第一个 token 作为主对象名
    - 提取结果若为单字符、参数键名（snake_case）、或占位词 -> 回退默认名，
      避免产物目录出现 "A" / "task_confirm" 这类垃圾名
    - 截断到 max_len（默认 40）
    """
    if not text:
        return _DEFAULT_OBJECT_NAME
    first = re.split(r"[,，\s\n]+", text.strip(), maxsplit=1)[0]
    name = safe_object_name(first, max_len=max_len)
    if not name or name in _JUNK_OBJECT_NAMES or len(name) <= 1:
        return _DEFAULT_OBJECT_NAME
    return name


class ArtifactService:
    def __init__(self, db: Database, artifact_root: str):
        self.db = db
        self.root = Path(artifact_root)

    def _dir_for(self, namespace: str) -> Path:
        d = (self.root / namespace).resolve()
        if not str(d).startswith(str(self.root.resolve())):
            raise NamespaceError("namespace 越界")
        return d

    def write(self, namespace: str, conv_id: int, agent_id: str, object_name: str,
              files: Dict[str, Union[str, bytes]], *, task_id: int | None = None,
              trace_id: str = "") -> Dict:
        d = self._dir_for(namespace)
        d.mkdir(parents=True, exist_ok=True)
        # 归属解析：显式 task_id 优先；否则从会话反向读取 空间/子系统/任务
        requested_task_id = task_id
        space_id = sub_system_id = task_id = None
        if conv_id is not None:
            from jarvis_agent.db import Conversation
            with self.db.session() as s:
                conv = s.get(Conversation, conv_id)
                if conv:
                    space_id, sub_system_id, task_id = conv.space_id, conv.sub_system_id, conv.task_id
        task_id = requested_task_id if requested_task_id is not None else task_id
        records = []
        for fname, content in files.items():
            if "/" in fname or "\\" in fname or fname in ("", ".", ".."):
                raise NamespaceError(f"非法文件名: {fname}")
            fpath = d / fname
            ftype = fpath.suffix.lstrip(".").lower() or "txt"
            version, parent_version = 1, None
            version_dir = d
            if task_id is not None:
                version, parent_version = self._next_version(
                    task_id=task_id, agent_id=agent_id,
                    object_name=object_name, file_type=ftype,
                )
                version_dir = d / f"v{version}"
                version_dir.mkdir(parents=True, exist_ok=True)
                fpath = version_dir / fname
            if isinstance(content, bytes):
                fpath.write_bytes(content)
            else:
                fpath.write_text(content, encoding="utf-8")
            with self.db.session() as s:
                rec = ArtifactRecord(namespace=namespace, conv_id=conv_id, agent_id=agent_id,
                                     object_name=object_name, file_type=ftype, path=str(fpath),
                                     space_id=space_id, sub_system_id=sub_system_id,
                                     task_id=task_id, version=version,
                                     parent_version=parent_version,
                                     trace_id=(trace_id or "").strip())
                s.add(rec)
                s.commit()
                s.refresh(rec)
                records.append(self._to_dict(rec))
        return {"namespace": namespace, "files": records}

    def _next_version(self, *, task_id: int, agent_id: str,
                      object_name: str, file_type: str) -> tuple[int, int | None]:
        versions = self._version_rows(task_id, agent_id, object_name, file_type)
        if not versions:
            return 1, None
        latest = versions[-1].version or 1
        return int(latest) + 1, int(latest)

    def _version_rows(self, task_id: int, agent_id: str,
                      object_name: str, file_type: str):
        with self.db.session() as s:
            return list(s.scalars(
                select(ArtifactRecord)
                .where(ArtifactRecord.task_id == task_id,
                       ArtifactRecord.agent_id == agent_id,
                       ArtifactRecord.object_name == object_name,
                       ArtifactRecord.file_type == file_type)
                .order_by(ArtifactRecord.id.asc())
            ))

    def list_namespaces(self, conv_id: Optional[int] = None) -> List[str]:
        with self.db.session() as s:
            q = select(ArtifactRecord.namespace).distinct()
            if conv_id is not None:
                q = q.where(ArtifactRecord.conv_id == conv_id)
            return sorted(set(s.scalars(q)))

    def list_files(self, namespace: str) -> List[Dict]:
        with self.db.session() as s:
            rows = list(s.scalars(select(ArtifactRecord).where(ArtifactRecord.namespace == namespace)))
            return [self._to_dict(r) for r in rows]

    def list_by_task(self, task_id: int) -> List[Dict]:
        with self.db.session() as s:
            rows = list(s.scalars(select(ArtifactRecord).where(ArtifactRecord.task_id == task_id)))
        return [self._to_dict(r) for r in rows]

    def read(self, record_id: int) -> Optional[str]:
        with self.db.session() as s:
            rec = s.get(ArtifactRecord, record_id)
            if not rec:
                return None
            p = Path(rec.path).resolve()
        if not str(p).startswith(str(self.root.resolve())):
            raise NamespaceError("产物路径越界")
        if p.suffix.lower() == ".xlsx":
            return _xlsx_preview(p)
        return p.read_text(encoding="utf-8")

    def read_bytes(self, record_id: int) -> Optional[bytes]:
        record = self._record(record_id)
        return self._read_record_bytes(record)

    def file_path(self, record_id: int) -> Path:
        record = self._record(record_id)
        path = Path(record.path).resolve()
        if not str(path).startswith(str(self.root.resolve())):
            raise NamespaceError("产物路径越界")
        return path

    def version_record(self, record_id: int, version: int):
        base = self._record(record_id)
        rows = self._version_rows(
            base.task_id, base.agent_id, base.object_name, base.file_type
        ) if base.task_id is not None else []
        for row in rows:
            if int(row.version or 1) == int(version):
                return row
        raise ValueError(f"产物版本不存在: {record_id} v{version}")

    def list_versions(self, record_id: int) -> List[Dict]:
        base = self._record(record_id)
        if base.task_id is None:
            return [self._to_dict(base)]
        return [self._to_dict(row) for row in self._version_rows(
            base.task_id, base.agent_id, base.object_name, base.file_type
        )]

    def read_version(self, record_id: int, version: int) -> Optional[str]:
        return self._read_record(self.version_record(record_id, version))

    def diff(self, record_id: int, left_version: int, right_version: int) -> Dict:
        left = self.version_record(record_id, left_version)
        right = self.version_record(record_id, right_version)
        left_lines = (self._read_record(left) or "").splitlines(keepends=True)
        right_lines = (self._read_record(right) or "").splitlines(keepends=True)
        matcher = difflib.SequenceMatcher(a=left_lines, b=right_lines, autojunk=False)
        blocks: List[Dict] = []
        summary = {"added": 0, "removed": 0, "changed": 0}
        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag == "equal":
                blocks.append({
                    "kind": "context",
                    "left_start": i1 + 1, "left_end": i2,
                    "right_start": j1 + 1, "right_end": j2,
                    "left": _display_lines(left_lines, i1, i2),
                    "right": _display_lines(right_lines, j1, j2),
                })
                continue
            if tag in ("delete", "replace"):
                summary["removed"] += i2 - i1
            if tag == "insert":
                summary["added"] += j2 - j1
            if tag == "replace":
                summary["changed"] += 1
                summary["added"] += j2 - j1
                blocks.append({
                    "kind": "changed", "left_start": i1 + 1, "left_end": i2,
                    "right_start": j1 + 1, "right_end": j2,
                    "left": _display_lines(left_lines, i1, i2),
                    "right": _display_lines(right_lines, j1, j2),
                })
                blocks.append({
                    "kind": "removed", "left_start": i1 + 1, "left_end": i2,
                    "right_start": j1 + 1, "right_end": j1,
                    "left": _display_lines(left_lines, i1, i2), "right": [],
                })
                blocks.append({
                    "kind": "added", "left_start": i1 + 1, "left_end": i1,
                    "right_start": j1 + 1, "right_end": j2,
                    "left": [], "right": _display_lines(right_lines, j1, j2),
                })
            else:
                kind = "added" if tag == "insert" else "removed"
                blocks.append({
                    "kind": kind,
                    "left_start": i1 + 1, "left_end": i2,
                    "right_start": j1 + 1, "right_end": j2,
                    "left": _display_lines(left_lines, i1, i2),
                    "right": _display_lines(right_lines, j1, j2),
                })
        return {
            "record_id": record_id,
            "left_version": int(left_version), "right_version": int(right_version),
            "summary": summary, "blocks": blocks,
        }

    def _record(self, record_id: int) -> ArtifactRecord:
        with self.db.session() as s:
            record = s.get(ArtifactRecord, record_id)
        if record is None:
            raise ValueError(f"产物不存在: {record_id}")
        return record

    def _read_record(self, record: ArtifactRecord) -> Optional[str]:
        path = Path(record.path).resolve()
        if not str(path).startswith(str(self.root.resolve())):
            raise NamespaceError("产物路径越界")
        if path.suffix.lower() == ".xlsx":
            return _xlsx_preview(path)
        return path.read_text(encoding="utf-8")

    def _read_record_bytes(self, record: ArtifactRecord) -> bytes:
        path = Path(record.path).resolve()
        if not str(path).startswith(str(self.root.resolve())):
            raise NamespaceError("产物路径越界")
        return path.read_bytes()

    @staticmethod
    def _to_dict(r: ArtifactRecord) -> Dict:
        return {"id": r.id, "namespace": r.namespace, "conv_id": r.conv_id, "agent_id": r.agent_id,
                "object_name": r.object_name, "type": r.file_type, "path": r.path,
                "space_id": r.space_id, "sub_system_id": r.sub_system_id, "task_id": r.task_id,
                "kb_doc_id": r.kb_doc_id,
                "version": int(getattr(r, "version", 1) or 1),
                "parent_version": getattr(r, "parent_version", None),
                "trace_id": getattr(r, "trace_id", "") or "",
                "created_at": r.created_at.isoformat() if r.created_at else None}


def _display_lines(lines: List[str], start: int, end: int) -> List[str]:
    return [line.rstrip("\r\n") for line in lines[start:end]]


def _xlsx_preview(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        blocks = []
        for sheet in workbook.worksheets:
            blocks.append(f"### {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) if value is not None else "" for value in row]
                blocks.append("| " + " | ".join(values) + " |")
        return "\n\n".join(blocks) or "（空 Excel）"
    finally:
        workbook.close()
