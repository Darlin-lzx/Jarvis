"""附件解析 / 上传接口 / kg 实体检索 / stream 引用注入 测试"""
import io

import pytest
from fastapi.testclient import TestClient

from db_loc_agent import app as app_module
from db_loc_agent.app import create_app
from db_loc_agent.attachments import extract_text


class FakeEmbedding:
    """确定性向量：无需网络"""

    def __call__(self, input):
        texts = input if isinstance(input, list) else [input]
        return [[1.0 if (len(t) + i) % 2 == 0 else -1.0 for i in range(8)] for t in texts]


def _client(tmp_path, with_seed=False):
    db_path = str(tmp_path / "t.db")
    application = create_app(
        db_path=db_path,
        artifact_root=str(tmp_path / "artifacts"),
        config_override={"llm": {"provider": "ollama", "model": "qwen2.5",
                                 "base_url": "http://127.0.0.1:1"},
                         "mysql_url": ""},  # 强制 SQLite 隔离，不污染生产 MySQL
        embedding_function=FakeEmbedding(),
    )
    if with_seed:
        from db_loc_agent.bootstrap import bootstrap
        from db_loc_agent.db import Database
        bootstrap(Database(db_path))
    return TestClient(application)


# ---------- extract_text 单测 ----------

def test_extract_text_txt_md_sql():
    assert extract_text("a.txt", "你好 hello".encode("utf-8")).strip() == "你好 hello"
    assert "# 标题" in extract_text("b.md", "# 标题\n内容".encode("utf-8"))
    assert "CREATE TABLE" in extract_text("c.sql", "CREATE TABLE T(ID INT);".encode("utf-8"))
    # gbk 回退解码
    assert extract_text("d.txt", "中文测试".encode("gbk")).strip() == "中文测试"


def test_extract_text_ddl_markdown():
    # ddl 当作 sql 解析
    assert "CREATE TABLE" in extract_text("t.ddl", "CREATE TABLE T(ID INT);".encode("utf-8"))
    assert "CREATE OR REPLACE" in extract_text("p.ddl", "CREATE OR REPLACE PROCEDURE P AS".encode("utf-8"))
    # markdown 当作 md 解析
    assert "# 标题" in extract_text("d.markdown", "# 标题\n正文".encode("utf-8"))
    assert "正文" in extract_text("d.markdown", "# 标题\n正文".encode("utf-8"))


def test_extract_text_xls_friendly_error():
    # xls 旧版二进制格式 → 友好中文提示（不必装 pandas），提示另存为 xlsx
    with pytest.raises(ValueError, match="另存为|xlsx"):
        extract_text("old.xls", b"\xd0\xcf\x11\xe0 dummy")


def test_extract_text_unsupported_raises():
    with pytest.raises(ValueError):
        extract_text("e.exe", b"xx")
    with pytest.raises(ValueError):
        extract_text("noext", b"xx")
    with pytest.raises(ValueError):
        extract_text("empty.txt", b"  \n  ")


def test_extract_text_docx():
    from docx import Document
    doc = Document()
    doc.add_paragraph("附件段落一")
    doc.add_paragraph("附件段落二")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A1"
    table.cell(0, 1).text = "B1"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("t.docx", buf.getvalue())
    assert "附件段落一" in text and "附件段落二" in text
    # 表格单元格文本按行合并
    assert "A1 | B1" in text


def test_extract_text_xlsx():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["列A", "列B"])
    ws.append([1, 2])
    buf = io.BytesIO()
    wb.save(buf)
    text = extract_text("t.xlsx", buf.getvalue())
    assert "列A\t列B" in text
    assert "1\t2" in text


# ---------- POST /api/chat/attachments ----------

def test_upload_attachments_api(tmp_path):
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("a.txt", "附件文本内容 abc".encode("utf-8"), "text/plain")),
                      ("files", ("b.md", "## 迁移规范\n内容".encode("utf-8"), "text/markdown"))])
    body = r.json()
    assert r.status_code == 200
    assert body["success"] is True
    data = body["data"]
    assert len(data) == 2
    assert all("attachment_id" in i and i["text_len"] > 0 for i in data)
    assert not any("error" in i for i in data)
    # 产物已写入 ArtifactRecord（右侧产物面板可见）
    ar = c.get("/api/artifacts").json()
    assert ar["success"] is True


def test_upload_attachments_response_fields(tmp_path):
    """上传成功项需返回 id（产物记录 id）与 object_name（文件名去后缀），并与 attachment_id 兼容"""
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("fields.txt", "字段验证正文".encode("utf-8"), "text/plain"))])
    body = r.json()
    assert body["success"] is True
    item = body["data"][0]
    assert item.get("id"), item
    assert item["object_name"] == "fields"  # safe_object_name(文件名去后缀)
    assert item["attachment_id"] == item["id"]
    assert item["text_len"] > 0


def test_upload_attachments_xls_error_item(tmp_path):
    """xls 上传：单文件返回 error 而非整体 500，success 仍为 true"""
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("ok.txt", "好的".encode("utf-8"), "text/plain")),
                      ("files", ("old.xls", b"\xd0\xcf\x11\xe0 dummy", "application/octet-stream"))])
    body = r.json()
    assert body["success"] is True
    ok = [i for i in body["data"] if "error" not in i]
    bad = [i for i in body["data"] if "error" in i]
    assert len(ok) == 1 and len(bad) == 1
    assert "xlsx" in bad[0]["error"]


def test_upload_attachments_ddl_md_api(tmp_path):
    """ddl / markdown 后缀上传成功并生成产物"""
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("schema.ddl", "CREATE TABLE T(ID INT);".encode("utf-8"), "text/plain")),
                      ("files", ("doc.markdown", "# 文档\n正文".encode("utf-8"), "text/markdown"))])
    body = r.json()
    assert body["success"] is True
    data = body["data"]
    assert len(data) == 2 and not any("error" in i for i in data)
    assert [i["object_name"] for i in data] == ["schema", "doc"]
    assert all(i.get("id") for i in data)


def test_upload_attachments_named_conv(tmp_path):
    c = _client(tmp_path)
    conv = c.post("/api/conversations", json={"title": "附件会话"}).json()["data"]
    r = c.post("/api/chat/attachments",
               files=[("files", ("notes.txt", "会话附件正文".encode("utf-8"), "text/plain"))],
               data={"conv_id": str(conv["id"])})
    body = r.json()
    assert body["success"] is True
    item = body["data"][0]
    # 按会话可查到产物记录（namespace 以 conv 为前缀）
    arts = c.get("/api/artifacts", params={"conv_id": conv["id"]}).json()["data"]
    assert any(a["id"] == item["attachment_id"] for a in arts)


def test_upload_attachments_partial_failure(tmp_path):
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("ok.txt", "好的".encode("utf-8"), "text/plain")),
                      ("files", ("bad.xyz", b"data", "application/octet-stream"))])
    body = r.json()
    assert body["success"] is True  # 单个失败不整体 500
    ok = [i for i in body["data"] if "error" not in i]
    bad = [i for i in body["data"] if "error" in i]
    assert len(ok) == 1 and len(bad) == 1
    assert "不支持的附件类型" in bad[0]["error"]


def test_upload_attachments_all_failed(tmp_path):
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("bad1.xyz", b"1", "application/octet-stream")),
                      ("files", ("bad2.xyz", b"2", "application/octet-stream"))])
    body = r.json()
    assert body["success"] is False
    assert body["data"] and all("error" in i for i in body["data"])
    assert "detail" in body


def test_upload_attachments_no_files(tmp_path):
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments")
    body = r.json()
    assert body["success"] is False
    assert "detail" in body


def test_upload_attachments_auto_import(tmp_path):
    from db_loc_agent.db import Database, Space
    db = Database(str(tmp_path / "t.db"))
    with db.session() as s:
        sp = Space(name="测试空间", description="", kb_ids="[]", members="[]")
        s.add(sp)
        s.commit()
        s.refresh(sp)
        space_id = sp.id
    c = _client(tmp_path)
    r = c.post("/api/chat/attachments",
               files=[("files", ("kb.txt",
                                 ("这是一段用于分块的文本。" + "填充内容。" * 40).encode("utf-8"),
                                 "text/plain"))],
               data={"space_id": str(space_id), "auto_import": "1"})
    body = r.json()
    assert body["success"] is True
    item = body["data"][0]
    assert item.get("kb_doc_id")  # 自动入库成功
    docs = c.get("/api/kb/documents", params={"space_id": str(space_id)}).json()["data"]
    assert any(d["id"] == item["kb_doc_id"] for d in docs)


# ---------- GET /api/kg/entities ----------

def test_kg_entities_api(tmp_path):
    from db_loc_agent.db import Database, KGEntity
    c = _client(tmp_path)
    assert c.get("/api/kg/entities").json()["success"] is True
    # 空库返回空列表
    assert c.get("/api/kg/entities").json()["data"] == []
    db = Database(str(tmp_path / "t.db"))
    with db.session() as s:
        s.add(KGEntity(name="USER_ORDER", etype="table"))
        s.add(KGEntity(name="达梦兼容", etype="concept"))
        s.commit()
    r = c.get("/api/kg/entities", params={"q": "USER"})
    assert r.json()["success"] is True
    data = r.json()["data"]
    assert any(d["name"] == "USER_ORDER" and d["etype"] == "table" for d in data)
    # q 为空返回最近 N 条（含中文实体）
    r2 = c.get("/api/kg/entities", params={"limit": 5})
    assert len(r2.json()["data"]) >= 2


# ---------- /api/chat/stream 引用注入（不回归） ----------

class CaptureLLM:
    """捕获注入 system 的 mock LLM（invoke/stream/parse 行为与 FakeLLM 一致）"""

    def __init__(self, reply="已按要求回答"):
        self.reply = reply
        self.systems = []
        self._called = 0

    def bind_tools(self, schemas, **kw):
        return self

    def _record(self, messages):
        for m in (messages or []):
            if getattr(m, "type", "") == "system":
                self.systems.append(getattr(m, "content", "") or "")

    def invoke(self, messages, **kw):
        from langchain_core.messages import AIMessage
        self._called += 1
        self._record(messages)
        return AIMessage(content=self.reply)

    def stream(self, messages, **kw):
        from langchain_core.messages import AIMessage
        self._record(messages)
        yield AIMessage(content=self.reply)

    def parse(self, text):
        return {"intent": "nl2sql", "agent_id": "nl2sql_qa", "object_name": "",
                "source_db": "", "target_db": "", "skill_id": "", "confidence": 0.95}


def test_stream_with_references_and_attachments(tmp_path, monkeypatch):
    llm = CaptureLLM()
    monkeypatch.setattr(app_module, "create_chat_model", lambda cfg: llm)
    c = _client(tmp_path, with_seed=True)
    conv = c.post("/api/conversations", json={"title": "引用测试"}).json()["data"]
    conv_id = conv["id"]
    up = c.post("/api/chat/attachments",
                files=[("files", ("notes.txt", "这是附件正文 123".encode("utf-8"), "text/plain"))],
                data={"conv_id": str(conv_id)}).json()["data"][0]
    payload = {
        "message": "请根据引用与附件回答",
        "conv_id": conv_id,
        "references": [
            {"type": "agent", "id": "nl2sql_qa", "name": "知识问答助手"},
            {"type": "kb_doc", "id": 999999, "name": "不存在的文档"},  # 应被跳过
        ],
        "attachment_ids": [up["attachment_id"], 999999],  # 999999 不存在应被跳过
    }
    r = c.post("/api/chat/stream", json=payload)
    assert r.status_code == 200
    # SSE 完整走完（done 事件），未报错
    assert '"type": "done"' in r.text or '"type":"done"' in r.text
    assert '"type": "error"' not in r.text
    # 附件文本已作为最高优先级上下文注入 system
    assert any("本次引用/附件上下文" in s and "这是附件正文 123" in s for s in llm.systems)
    # agent 引用注入「名称：描述」
    assert any("知识问答助手" in s for s in llm.systems)
    # 不存在的引用未进入上下文，也未阻塞
    assert not any("不存在的文档" in s for s in llm.systems)


class OrchLLM(CaptureLLM):
    """编排测试：第 1 次 invoke 返回 call_agent 委派工具调用，之后返回最终答复"""

    def __init__(self, reply="已完成"):
        super().__init__(reply=reply)
        self._calls = 0

    def invoke(self, messages, **kw):
        from langchain_core.messages import AIMessage
        self._calls += 1
        self._record(messages)
        if self._calls == 1:
            return AIMessage(content="需要委派给转换助手", tool_calls=[
                {"name": "call_agent",
                 "args": {"agent_id": "ddl_converter", "question": "请转换 select * from t"},
                 "id": "call_agent_1"}])
        return AIMessage(content=self.reply)

    def stream(self, messages, **kw):
        from langchain_core.messages import AIMessage
        self._calls += 1
        self._record(messages)
        if self._calls == 1:
            yield AIMessage(content="需要委派给转换助手", tool_calls=[
                {"name": "call_agent",
                 "args": {"agent_id": "ddl_converter", "question": "请转换 select * from t"},
                 "id": "call_agent_1"}])
            return
        yield AIMessage(content=self.reply)


def test_stream_sub_agent_and_system_events(tmp_path, monkeypatch):
    """主 Agent 委派 call_agent 时：SSE 流实时下发「系统调度事件 + 子 Agent 输出事件」，
    且子 Agent 带调度溯源（dispatcher_name）"""
    llm = OrchLLM()
    monkeypatch.setattr(app_module, "create_chat_model", lambda cfg: llm)
    c = _client(tmp_path, with_seed=True)
    r = c.post("/api/chat/stream", json={"message": "帮我做个 DDL 转换"})
    assert r.status_code == 200
    # 完整走完，不报错
    assert '"type": "done"' in r.text or '"type":"done"' in r.text
    assert '"type": "error"' not in r.text
    # 系统调度事件：正在调用子 Agent
    assert '"type": "system"' in r.text or '"type":"system"' in r.text
    assert "正在调用【小转】执行任务…" in r.text
    # 子 Agent 输出事件：名称 + 回答 + 调度溯源（派遣方为当前主 Agent 知识问答助手）
    assert '"type": "sub_agent"' in r.text or '"type":"sub_agent"' in r.text
    assert '"agent_name": "小转"' in r.text
    assert '"reply": "已完成"' in r.text
    assert '"dispatcher_name": "小问"' in r.text


def test_stream_without_extra_fields(tmp_path, monkeypatch):
    """不带 references/attachment_ids 的既有调用完全兼容"""
    llm = CaptureLLM()
    monkeypatch.setattr(app_module, "create_chat_model", lambda cfg: llm)
    c = _client(tmp_path)
    r = c.post("/api/chat/stream", json={"message": "普通问答"})
    assert r.status_code == 200
    assert '"type": "done"' in r.text or '"type":"done"' in r.text
    assert not any("本次引用" in s for s in llm.systems)