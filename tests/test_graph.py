import json
import threading
from pathlib import Path

from db_loc_agent.agents import AgentService
from db_loc_agent.artifacts import ArtifactService
from db_loc_agent.db import Database
from db_loc_agent.graph import AgentState, build_graph
from db_loc_agent.graph import GraphRuntime
from db_loc_agent.config import AppConfig
from db_loc_agent.prompts import PromptService
from db_loc_agent.skills import SkillService
from langchain_core.language_models.fake_chat_models import GenericFakeChatModel
from langchain_core.messages import AIMessage


class FakeLLM:
    """非流式 mock：invoke 返回 AIMessage；with_structured_output 由外层独立实现，此处辅助 run 使用"""

    def __init__(self, reply="转换完成", parse=None, tool_script=None, skill_reply="脚本输出 100GB=102400MB"):
        self.reply = reply
        self._parse = parse
        # 模拟：首次调用绑定工具后返回 use_skill 工具调用，随后返回最终回复
        self._tool_script = tool_script
        self._skill_reply = skill_reply
        self._called = 0
        self._schemas = []

    def bind_tools(self, schemas, **kw):
        self._schemas = schemas
        return self

    def invoke(self, messages, **kw):
        from langchain_core.messages import AIMessage
        self._called += 1
        if self._tool_script and self._called == 1:
            return AIMessage(content="我需要调用技能", tool_calls=[
                {"name": "use_skill", "args": {"skill_id": self._tool_script, "args": "10"},
                 "id": "call_1"}])
        return AIMessage(content=self.reply)

    def stream(self, messages, **kw):
        """最终回复轮流式重放（graph.execute 对无工具调用轮走 stream 分支）"""
        from langchain_core.messages import AIMessage
        yield AIMessage(content=self.reply)

    def parse(self, text):
        return self._parse or {"intent": "ddl_convert", "agent_id": "ddl_converter",
                               "object_name": "user_order", "source_db": "oracle",
                               "target_db": "dm", "skill_id": "", "confidence": 0.95}


class ScriptedLLM(FakeLLM):
    """按调用次数依次返回预设内容/tool_calls（用于编排类测试）"""

    def __init__(self, steps, parse=None):
        super().__init__(reply="", parse=parse)
        self.steps = steps
        self.calls = []

    def invoke(self, messages, **kw):
        from langchain_core.messages import AIMessage
        idx = len(self.calls)
        if idx >= len(self.steps):
            return AIMessage(content="（无更多预设步骤）")
        text, tool_calls = self.steps[idx]
        self.calls.append(text if isinstance(text, str) else text)
        return AIMessage(content=text if isinstance(text, str) else "",
                         tool_calls=tool_calls or [])

    def stream(self, messages, **kw):
        """流式按调用次数返回当前步骤（execute 主循环统一用 stream，含工具调用轮）。"""
        from langchain_core.messages import AIMessage
        idx = len(self.calls)
        if idx >= len(self.steps):
            yield AIMessage(content="（无更多预设步骤）")
            return
        text, tool_calls = self.steps[idx]
        self.calls.append(text if isinstance(text, str) else text)
        yield AIMessage(content=text if isinstance(text, str) else "",
                        tool_calls=tool_calls or [])


def _runtime_for_routing():
    return GraphRuntime(
        llm=None,
        rag_search=lambda *args, **kwargs: [],
        prompt_service=None,
        artifact_service=None,
        skill_service=None,
        parse_schema={},
        config=AppConfig(),
    )


def test_graph_full_flow(tmp_path):
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    prompt_svc.create("ddl_converter 提示词", "agent",
                      AgentService(db).get("ddl_converter")["default_prompt"])

    graph = build_graph(
        llm=FakeLLM(reply="```sql\nCREATE TABLE USER_ORDER (ID INT);\n```"),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc,
        agent_service=AgentService(db),
        artifact_service=artifact_svc,
        skill_service=None,
        parse_schema={},
        max_iterations=1,
    )
    state = AgentState(user_input="将 user_order 表转成达梦 DDL", conv_id=1)
    final = graph.invoke(state)
    assert final["reply"]
    assert final["namespace"].startswith("sess_1/ddl_converter/user_order/")
    namespaces = artifact_svc.list_namespaces(conv_id=1)
    assert len(namespaces) == 1


def test_graph_llm_stream_stops_after_user_cancel(tmp_path):
    """会话中止事件必须在主 LLM 流的下一个 chunk 边界终止图执行。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    cancel_event = threading.Event()
    cancel_event.set()

    class CancelledStreamLLM(FakeLLM):
        def stream(self, messages, **kwargs):
            from langchain_core.messages import AIMessage
            yield AIMessage(content="第一段")
            yield AIMessage(content="这段不应被处理")

    graph = build_graph(
        llm=CancelledStreamLLM(reply="", parse={
            "intent": "ddl_convert", "agent_id": "ddl_converter",
            "object_name": "user_order", "source_db": "oracle",
            "target_db": "dm", "skill_id": "", "confidence": 0.95,
        }),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=PromptService(db),
        agent_service=AgentService(db),
        artifact_service=artifact_svc,
        skill_service=None,
        parse_schema={},
        max_iterations=1,
        cancellation_event=cancel_event,
    )

    import pytest
    with pytest.raises(RuntimeError, match="用户已中止"):
        list(graph.stream(AgentState(user_input="转换", conv_id=1),
                          stream_mode=["updates", "custom"]))


def test_graph_low_confidence_fallback(tmp_path):
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    graph = build_graph(
        llm=FakeLLM(reply="回答", parse={"intent": 999, "agent_id": "", "object_name": "x",
                                         "source_db": "", "target_db": "", "skill_id": "",
                                         "confidence": 0.2}),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=PromptService(db),
        agent_service=AgentService(db),
        artifact_service=artifact_svc,
        skill_service=None,
        parse_schema={},
        max_iterations=1,
    )
    final = graph.invoke(AgentState(user_input="随便问问", conv_id=2))
    assert final["agent_id"] == "nl2sql_qa"  # 回退知识库问答


def test_agent_autonomously_calls_skill(tmp_path):
    """agent 自主调用 use_skill 工具：执行技能脚本并把输出注入上下文"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    skill_svc = SkillService(db, str(tmp_path / "skills"))
    text = """---
name: disk_conv
description: GB 转 MB 换算
script: conv.py
---

按脚本换算容量单位。

```python
import sys
v = float(sys.argv[1]) if len(sys.argv) > 1 else 100
print(f"{v}GB = {v*1024}MB")
```
"""
    skill = skill_svc.create_from_md(text)
    skill_svc.set_enabled(skill["skill_id"], True)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    prompt_svc = PromptService(db)
    prompt_svc.create("ddl_converter 提示词", "agent",
                      AgentService(db).get("ddl_converter")["default_prompt"])
    # 定义解析结果不要带 skill_id（用户未 @），看 agent 是否自主调用
    graph = build_graph(
        llm=FakeLLM(reply="换算结果：10GB = 10240MB", tool_script=skill["skill_id"]),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc,
        agent_service=AgentService(db),
        artifact_service=artifact_svc,
        skill_service=skill_svc,
        parse_schema={},
        max_iterations=3,
    )
    state = AgentState(user_input="把 10GB 换算成 MB", conv_id=3)
    final = graph.invoke(state)
    assert final["reply"] == "换算结果：10GB = 10240MB"
    # 工具调用已执行：技能脚本被执行（脚本真实落盘运行）
    r = skill_svc.run_script(skill["skill_id"], args="10")
    assert "10240" in r["stdout"]


def test_multi_agent_orchestration_call_agent(tmp_path):
    """多 Agent 编排：主 agent 委派 call_agent 给另一个 agent；子 agent 不带 call_agent 工具（防自递归）"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    prompt_svc.create("ddl_converter 提示词", "agent",
                      AgentService(db).get("ddl_converter")["default_prompt"])
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    call_tool = {"name": "call_agent",
                 "args": {"agent_id": "ddl_converter",
                          "question": "请转换 select * from t"},
                 "id": "call_agent_1"}
    llm = ScriptedLLM([
        ("需要委派给转换智能体", [call_tool]),
        ("子智能体回复：已转换 SQL", None),
        ("最终汇总答复：已由 DDL 转换智能体完成", None),
    ])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(
        user_input="帮我做个 DDL 转换", conv_id=4, agent_id="ddl_converter"))
    assert "最终汇总答复" in final["reply"]
    assert final.get("agent_calls") and final["agent_calls"][0]["agent_id"] == "ddl_converter"
    assert final["agent_calls"][0]["reply"] == "子智能体回复：已转换 SQL"
    # 主图暴露过 call_agent 工具；子 agent 不允许再委派（递归防护）
    main_schemas = {s["function"]["name"] for s in llm._schemas} if llm._schemas else set()
    assert "call_agent" in main_schemas


def test_custom_stream_events_on_call_agent(tmp_path):
    """call_agent 委派时经 custom 流模式实时下发「系统调度事件 + 子 Agent 输出事件」"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    prompt_svc.create("ddl_converter 提示词", "agent",
                      AgentService(db).get("ddl_converter")["default_prompt"])
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    call_tool = {"name": "call_agent",
                 "args": {"agent_id": "ddl_converter",
                          "question": "请转换 select * from t"},
                 "id": "call_agent_1"}
    llm = ScriptedLLM([
        ("需要委派给转换智能体", [call_tool]),
        ("子智能体回复：已转换 SQL", None),
        ("最终汇总答复：已由 SQL 转换助手完成", None),
    ])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for _mode, payload in graph.stream(
            AgentState(user_input="帮我做个 DDL 转换", conv_id=5,
                       agent_id="ddl_converter"),
            stream_mode=["custom"]):
        custom_events.append(payload)
    kinds = [e.get("type") for e in custom_events if isinstance(e, dict)]
    assert "system" in kinds and "sub_agent" in kinds
    call_events = [e for e in custom_events if isinstance(e, dict)
                   and e.get("type") == "system"
                   and e.get("event") in ("calling_agent", "agent_finished")]
    assert [(e.get("event"), e.get("agent_id")) for e in call_events] == [
        ("calling_agent", "ddl_converter"),
        ("agent_finished", "ddl_converter"),
    ]
    assert "执行完成" in call_events[-1]["message"]
    sys_ev = next(e for e in custom_events if isinstance(e, dict) and e.get("type") == "system")
    sub_ev = next(e for e in custom_events if isinstance(e, dict) and e.get("type") == "sub_agent")
    assert "正在调用" in sys_ev["message"] and "小转" in sys_ev["message"]
    assert sub_ev["agent_name"] == "小转"
    assert sub_ev["reply"] == "子智能体回复：已转换 SQL"
    assert sub_ev["dispatcher_name"]  # 主 Agent 溯源名非空


def test_call_agent_tool_message_omits_displayed_sub_reply(tmp_path):
    """子 Agent 完整回复已单独展示给用户，返回主 Agent 的 ToolMessage 不应携带正文。

    否则指挥官会把 ToolMessage 中的完整报告复述进最终回复，前端出现两份相同内容。
    """
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    prompt_svc.create("ddl_converter 提示词", "agent",
                      AgentService(db).get("ddl_converter")["default_prompt"])
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    call_tool = {"name": "call_agent",
                 "args": {"agent_id": "ddl_converter",
                          "question": "请转换 select * from t"},
                 "id": "call_agent_1"}
    sub_reply = "子智能体完整回复：" + "这是一份很长的转换报告内容。" * 30

    class RecordingScriptedLLM(ScriptedLLM):
        def __init__(self, steps):
            super().__init__(steps)
            self.seen_messages = []

        def stream(self, messages, **kw):
            self.seen_messages.append(list(messages))
            yield from super().stream(messages, **kw)

    llm = RecordingScriptedLLM([
        ("需要委派给转换智能体", [call_tool]),
        (sub_reply, None),
        ("转换已完成，详见上方子智能体输出。", None),
    ])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(
        user_input="帮我做个 DDL 转换", conv_id=6, agent_id="ddl_converter"))

    tool_payloads = []
    for messages in llm.seen_messages:
        for msg in messages:
            if msg.__class__.__name__ == "ToolMessage" and msg.name == "call_agent":
                tool_payloads.append(json.loads(msg.content))
    assert len(tool_payloads) == 1
    assert sub_reply not in tool_payloads[0].get("reply", "")
    assert tool_payloads[0]["reply_already_displayed"] is True
    assert final["agent_calls"][0]["reply"] == sub_reply


def test_call_agent_nested_tokens_are_not_parent_stream(tmp_path):
    """子 Agent 内部 LLM token 不能进入父图 messages 流。

    app.py 会把 messages 流当主回复 text 下发；若嵌套子 Agent token 也被
    LangGraph 捕获，前端会同时收到 text 与 sub_agent 两份相同内容。
    """
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    prompt_svc.create("ddl_converter 提示词", "agent",
                      AgentService(db).get("ddl_converter")["default_prompt"])
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    class CallbackToolFake(GenericFakeChatModel):
        def _stream(self, messages, stop=None, run_manager=None, **kw):
            from langchain_core.messages import AIMessageChunk
            from langchain_core.outputs import ChatGenerationChunk
            result = self._generate(messages, stop=stop,
                                    run_manager=run_manager, **kw)
            msg = result.generations[0].message
            if msg.content:
                chunk_msg = AIMessageChunk(content=msg.content, chunk_position="last")
                if run_manager:
                    run_manager.on_llm_new_token(msg.content, chunk=chunk_msg)
                yield ChatGenerationChunk(message=chunk_msg)
            if getattr(msg, "tool_calls", None):
                chunks = []
                for i, tc in enumerate(msg.tool_calls):
                    chunks.append({
                        "name": tc["name"],
                        "args": json.dumps(tc["args"], ensure_ascii=False),
                        "id": tc["id"], "index": i, "type": "tool_call_chunk",
                    })
                yield ChatGenerationChunk(message=AIMessageChunk(
                    content="", tool_call_chunks=chunks,
                    chunk_position="last" if not msg.content else None))

    sub_reply = "子智能体流式完整回复：这是详细报告"
    responses = iter([
        json.dumps({"intent": "migration_overall", "agent_id": "migration_commander",
                    "object_name": "", "source_db": "", "target_db": "",
                    "skill_id": "", "params": {}, "confidence": 1.0},
                   ensure_ascii=False),
        AIMessage(content="我来调度", tool_calls=[{
            "name": "call_agent",
            "args": {"agent_id": "ddl_converter", "question": "请转换"},
            "id": "nested_call_1",
        }]),
        sub_reply,
        "主智能体简要总结",
    ])
    graph = build_graph(
        llm=CallbackToolFake(messages=responses),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    parent_texts = []
    final_reply = ""
    for mode, payload in graph.stream(
            AgentState(user_input="帮我做个 DDL 转换", conv_id=7,
                       agent_id="ddl_converter"),
            stream_mode=["messages", "updates"]):
        if mode == "messages":
            parent_texts.append(payload[0].content or "")
        elif mode == "updates" and isinstance(payload, dict) and payload:
            update = next(iter(payload.values()), {})
            if isinstance(update, dict) and "reply" in update:
                final_reply = update.get("reply", "")

    assert sub_reply not in "".join(parent_texts)
    assert final_reply == "主智能体简要总结"


def test_parallel_calls_cannot_start_future_migration_stages(tmp_path):
    """逐级确认模式下，条件边只放行当前阶段智能体。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    llm = ScriptedLLM([("方案专家完整方案", None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=6,
        agent_service=AgentService(db),
    )
    custom_events = []
    final = {}
    for mode, payload in graph.stream(
            AgentState(
                user_input="继续生成迁移方案", conv_id=8,
                agent_id="migration_commander",
                task_state={"current_step": "SOLUTION_GENERATE",
                            "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict) \
                and payload.get("event") == "calling_agent":
            custom_events.append(payload["agent_id"])
        elif mode == "updates" and isinstance(payload, dict) and payload:
            update = next(iter(payload.values()), {})
            if isinstance(update, dict) and "agent_calls" in update:
                final = update

    assert custom_events == ["migration_planner"]
    assert [c["agent_id"] for c in final["agent_calls"]] == ["migration_planner"]


def test_sub_agent_conversational_reply_does_not_create_artifact(tmp_path):
    """子 Agent 的追问/中间对话不是交付物，不得写入右侧产物面板。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    call_tool = {"name": "call_agent",
                 "args": {"agent_id": "migration_assessor",
                          "question": "请评估迁移风险"},
                 "id": "call_conversational_1"}
    llm = ScriptedLLM([
        ("需要评估", [call_tool]),
        ("请先提供源库对象清单、目标库类型和停机窗口。", None),
        ("评估专家需要补充信息。", None),
    ])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(
        user_input="继续评估", conv_id=9,
        agent_id="migration_commander",
        task_state={"current_step": "ASSESSMENT_RUNNING",
                    "step_confirm": True, "pipeline_start": True}))

    assert final["agent_calls"][0]["artifacts"] == []
    assert artifact_svc.list_namespaces(conv_id=9) == []


def test_sub_agent_core_report_creates_one_document_and_chat_link(tmp_path):
    """核心报告只落业务文档；聊天事件只给简短状态和可预览文件记录。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    report = "# 迁移评估报告\n\n源库与目标库兼容性良好，主要风险集中在函数兼容性。"
    llm = ScriptedLLM([(report, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="继续评估", conv_id=10,
                       agent_id="migration_commander",
                       task_state={"current_step": "ASSESSMENT_RUNNING",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    files = event["artifacts"]
    assert [f["path"].split("/")[-1] for f in files] == ["评估报告.md"]
    assert report not in event["reply"]
    assert "评估报告.md" in event["reply"]
    assert artifact_svc.read(files[0]["id"]) == report


def test_collector_failure_report_creates_chinese_artifact(tmp_path):
    """小采的结构化采集失败报告仍是核心产物，且使用中文文件名。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    report = (
        "# Oracle 源库数据采集报告\n\n"
        "## 一、采集状态：无法完成\n\n"
        "| 数据源名称 | 类型 | 连接状态 |\n"
        "|:---|:---|:---:|\n"
        "| src_o | Oracle | 失败 |\n\n"
        "数据源未配置，无法采集对象清单。"
    )
    llm = ScriptedLLM([(report, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="数据迁移", conv_id=11,
                       agent_id="migration_commander",
                       task_state={"current_step": "DATA_COLLECTION",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    files = event["artifacts"]
    assert [f["path"].split("/")[-1] for f in files] == ["数据采集报告.md"]
    assert report not in event["reply"]
    assert "数据采集报告.md" in event["reply"]
    assert artifact_svc.read(files[0]["id"]) == report


def test_sub_agent_confirmation_form_is_not_written_as_artifact(tmp_path):
    """关键信息确认不是核心交付物，UI 控制块不得写入右侧产物文件。"""
    from db_loc_agent.graph import _is_core_delivery
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    form = {
        "title": "迁移方案关键信息确认",
        "fields": [{"name": "source_db_type", "label": "源数据库类型",
                    "type": "text", "options": []}],
        "submit": "确认并生成方案",
    }
    assert not _is_core_delivery(
        "migration_planner",
        "## 当前评估状态\n\n- 待补充：源数据库类型。\n\n"
        "暂无法输出兼容性分析。请通过下方表单补充关键信息。",
    )
    confirmation = (
        "## 当前评估状态\n\n"
        "- ⏳ 待补充：源数据库类型\n\n"
        "当前评估状态：待补充。暂无法输出兼容性分析。请通过下方表单补充关键信息：\n\n"
        "<UI_RENDER>\n" + json.dumps(form, ensure_ascii=False) + "\n</UI_RENDER>\n\n"
        "提供以上信息后，我将生成可执行方案。"
    )
    llm = ScriptedLLM([(confirmation, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    final = {}
    for mode, payload in graph.stream(
            AgentState(user_input="继续", conv_id=14,
                       agent_id="migration_commander",
                       task_state={"current_step": "SOLUTION_GENERATE",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)
        elif mode == "updates" and isinstance(payload, dict) and payload:
            update = next(iter(payload.values()), {})
            if isinstance(update, dict) and "agent_calls" in update:
                final = update

    event = custom_events[0]
    assert event["artifacts"] == []
    assert "<UI_RENDER>" not in event["reply"]
    assert event["ui_render"] == form
    assert final["agent_calls"][0]["artifacts"] == []
    assert artifact_svc.list_namespaces(conv_id=14) == []


def test_migration_plan_artifact_strips_interactive_appendix(tmp_path):
    """小方核心方案不携带 UI_RENDER 附录或数据源选择交互内容。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    reply = (
        "# MySQL 到达梦迁移落地方案\n\n"
        "## 一、实施路径\n\n"
        "1. 完成对象清点与兼容性评估。\n"
        "2. 分批执行结构迁移与数据校验。\n\n"
        "## 二、风险控制\n\n"
        "- 保留回滚窗口，迁移前后执行数据校验。\n\n"
        "## 附录UI_RENDER\n\n"
        "请选择数据源：源数据源、目标数据源。\n\n"
        "UI_RENDER: {\"title\":\"选择数据源\",\"fields\":[{\"name\":\"datasource\"}]}\n"
    )
    llm = ScriptedLLM([(reply, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="生成迁移方案", conv_id=15,
                       agent_id="migration_commander",
                       task_state={"current_step": "SOLUTION_GENERATE",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    assert [f["path"].split("/")[-1] for f in event["artifacts"]] == ["迁移方案.md"]
    content = artifact_svc.read(event["artifacts"][0]["id"])
    assert "MySQL 到达梦迁移落地方案" in content
    assert "UI_RENDER" not in content
    assert "选择数据源" not in content


def test_substantive_migration_plan_with_pending_risk_still_creates_artifact(tmp_path):
    """完整迁移方案不因正文出现“待确认”风险描述而被误判为非交付。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    reply = (
        "# MySQL 数据库迁移方案\n\n"
        "## 一、实施路径\n\n"
        "1. 先完成结构迁移，再执行全量数据迁移。\n"
        "2. 迁移完成后执行行数校验与核心业务抽样比对。\n\n"
        "## 二、风险控制\n\n"
        "- 割接窗口和回滚预案需在实施前待确认。\n"
        "- 目标库索引需在数据导入后重建并验证执行计划。\n"
    )
    llm = ScriptedLLM([(reply, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="生成迁移方案", conv_id=19,
                       agent_id="migration_commander",
                       task_state={"current_step": "SOLUTION_GENERATE",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    assert [f["path"].split("/")[-1] for f in event["artifacts"]] == ["迁移方案.md"]
    content = artifact_svc.read(event["artifacts"][0]["id"])
    assert "MySQL 数据库迁移方案" in content
    assert "实施路径" in content
    assert "待确认" in content


def test_code_migrator_builds_sql_migration_report_artifacts(tmp_path):
    """小改生成改造前后 SQL 与目标库模拟执行结果，不落盘交互表单。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    reply = (
        "## 改造前 SQL\n\n"
        "```sql\nSELECT * FROM user_order WHERE ROWNUM <= 100;\n```\n\n"
        "## 改造后 SQL\n\n"
        "```sql\nSELECT * FROM user_order LIMIT 100;\n```\n\n"
        "## 风险说明\n\n"
        "- LIMIT 语法依赖目标库 MySQL 兼容模式。\n\n"
        "UI_RENDER: {\"title\":\"是否需要进一步优化？\",\"fields\":[]}\n"
    )
    llm = ScriptedLLM([(reply, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="改造订单查询 SQL", conv_id=16,
                       agent_id="migration_commander",
                       task_state={"current_step": "CODE_MIGRATE",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    filenames = [f["path"].split("/")[-1] for f in event["artifacts"]]
    assert filenames == ["改造后SQL.sql", "SQL改造报告.md", "代码改造清单.xlsx"]
    report = artifact_svc.read(event["artifacts"][1]["id"])
    converted = artifact_svc.read(event["artifacts"][0]["id"])
    assert "SELECT * FROM user_order WHERE ROWNUM <= 100;" in report
    assert "SELECT * FROM user_order LIMIT 100;" in report
    assert "执行模式：SIMULATION" in report
    assert "模拟执行成功" in report
    assert "UI_RENDER" not in report
    assert converted == "SELECT * FROM user_order LIMIT 100;"
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(artifact_svc.read_bytes(event["artifacts"][2]["id"])))
    sheet = workbook["改造清单"]
    assert [cell.value for cell in sheet[1]] == [
        "序号", "对象/代码类型", "改造前SQL", "改造后SQL", "目标库", "模拟执行结果", "风险说明"
    ]
    assert sheet["A2"].value == 1
    assert "ROWNUM <= 100" in sheet["C2"].value
    assert "LIMIT 100" in sheet["D2"].value


def test_code_migrator_structured_analysis_creates_report_artifact(tmp_path):
    """小改无 SQL 的结构化分析仍是交付物，且不落盘输出要求和追问。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    reply = (
        "【分析结果】\n\n"
        "数据源类型：Excel\n"
        "目标数据库类型：MySQL\n"
        "对象数量：3\n"
        "JDBC适配器：可能需要转换\n"
        "数据转换：可能需要处理\n"
        "事务处理：可能需要调整\n"
        "分页方法：可能需要调整\n\n"
        "【风险】\n\n"
        "数据连接问题\n"
        "存储过程调用问题\n"
        "数据转换不一致\n"
        "事务处理失败\n"
        "分页逻辑错误\n\n"
        "【输出文件】\n\n"
        "修改后的JDBC连接字符串\n"
        "调整后的存储过程调用方式\n"
        "数据转换处理脚本\n\n"
        "【下一步建议】\n\n"
        "验证和调整JDBC连接字符串，确保符合MySQL语法\n"
        "转换存储过程调用方式，确保语法正确\n"
        "请确认以上分析是否正确，并继续提供必要的数据转换规则或具体的JDBC连接参数。"
    )
    class RecordingScriptedLLM(ScriptedLLM):
        def __init__(self, steps):
            super().__init__(steps)
            self.seen_messages = []

        def stream(self, messages, **kw):
            self.seen_messages.append(list(messages))
            yield from super().stream(messages, **kw)

    llm = RecordingScriptedLLM([(reply, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="执行代码改造", conv_id=22,
                       agent_id="migration_commander",
                       task_state={"current_step": "CODE_MIGRATE",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    assert [f["path"].split("/")[-1] for f in event["artifacts"]] == ["代码改造分析报告.md"]
    report = artifact_svc.read(event["artifacts"][0]["id"])
    assert "【分析结果】" in report
    assert "【风险】" in report
    assert "【下一步建议】" in report
    assert "输出文件" not in report
    assert "请确认以上分析" not in report
    assert "JDBC连接参数" not in report
    code_system_prompts = [
        messages for messages in llm.seen_messages
        if messages and "代码改造助手" in str(messages[0].content)
    ]
    assert code_system_prompts
    system_prompt = str(code_system_prompts[0][0].content)
    assert "缺少待改造代码" in system_prompt
    assert "示例数据" in system_prompt
    assert "能力介绍或材料清单" in system_prompt
    assert "逐项引用资产名和原始SQL" in system_prompt
    assert "不得复述【输出文件】要求" in system_prompt


def test_code_migrator_missing_input_falls_back_to_mock_artifacts(tmp_path):
    """代码改造阶段缺材料时使用示例数据交付，能力说明不得替代产物。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    call_tool = {"name": "call_agent",
                 "args": {"agent_id": "code_migrator",
                          "question": "执行代码改造"},
                 "id": "call_code_missing_1"}
    llm = ScriptedLLM([
        ("调用小改", [call_tool]),
        ("## 我能做什么\n\n- 分析 SQL 方言差异\n\n"
         "## 需要你提供的材料\n\n- 待改造代码片段", None),
        ("小改执行完成。", None),
    ])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    final = {}
    for mode, payload in graph.stream(
            AgentState(user_input="执行小改", conv_id=23,
                       agent_id="migration_commander",
                       task_state={"current_step": "CODE_MIGRATE",
                                   "task_params": {"source_db_type": "Oracle",
                                                   "target_db_type": "MySQL"},
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)
        elif mode == "updates" and isinstance(payload, dict) and payload:
            update = next(iter(payload.values()), {})
            if isinstance(update, dict) and "agent_calls" in update:
                final = update

    filenames = [f["path"].split("/")[-1] for f in final["artifacts"]]
    assert filenames == ["改造后SQL.sql", "SQL改造报告.md", "代码改造清单.xlsx"]
    report = artifact_svc.read(final["artifacts"][1]["id"])
    assert "示例数据" in report
    assert "ROWNUM" in report
    assert "LIMIT" in report
    assert "执行模式：SIMULATION" in report
    assert "模拟执行成功" in report
    assert "我能做什么" not in report
    assert custom_events[0]["artifacts"]
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(artifact_svc.read_bytes(final["artifacts"][2]["id"])))
    sheet = workbook["改造清单"]
    assert sheet.max_row == 7
    categories = [sheet.cell(row=row, column=2).value for row in range(2, sheet.max_row + 1)]
    assert "SQL-分页查询" in categories
    assert "SQL-日期函数" in categories
    assert "SQL-存储过程调用" in categories
def test_verifier_builds_report_from_structured_assets_not_requirements(tmp_path):
    """小测产物必须是逐条 SQL 校验结果，不能把输出要求当作验证报告。"""
    from db_loc_agent.assets import MigrateAssetService

    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    migrate_assets = MigrateAssetService(db)
    migrate_assets.upsert(50, {
        "asset_type": "table", "asset_name": "users", "risk_level": "低",
        "original_sql": "CREATE TABLE users (id NUMBER)",
        "converted_sql": "CREATE TABLE users (id INT)",
        "suggestion": "已按 MySQL 类型改写 NUMBER。",
    })
    migrate_assets.upsert(50, {
        "asset_type": "view", "asset_name": "user_orders", "risk_level": "高",
        "original_sql": "CREATE OR REPLACE VIEW user_orders AS SELECT * FROM users",
        "converted_sql": "",
        "suggestion": "补充 MySQL 视图创建语句。",
    })

    requirement_reply = (
        "【分析结果】\n"
        "1. 数据源信息：Excel，目标库 MySQL，对象数量 3。\n"
        "【风险】\n"
        "1. 数据转换可能引入语法错误。\n"
        "【输出文件】\n"
        "1. 校验报告：包含成功和失败的SQL语句清单、原因分析及修复建议。\n"
        "【下一步建议】\n"
        "1. 确保目标库已创建。\n"
    )
    llm = ScriptedLLM([(requirement_reply, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db), migrate_assets=migrate_assets,
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="执行迁移验证", conv_id=17,
                       agent_id="migration_commander",
                       task_state={"current_step": "VERIFY_STAGE",
                                   "step_confirm": True, "pipeline_start": True,
                                   "migrate_task_id": 50,
                                   "task_params": {"target_db_type": "MySQL"}}),
            stream_mode=["custom"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    assert [f["path"].split("/")[-1] for f in event["artifacts"]] == ["校验报告.md", "迁移测试清单.xlsx"]
    report = artifact_svc.read(event["artifacts"][0]["id"])
    assert "目标库：MySQL" in report
    assert "校验对象总数：10" in report
    assert "示例数据" in report
    assert "成功：5" in report
    assert "失败：5" in report
    assert "`users`" in report
    assert "`user_orders`" in report
    assert "未生成转换 SQL" in report
    assert "补充 MySQL 视图创建语句" in report
    assert "CREATE TABLE users (id INT)" in report
    assert "【输出文件】" not in report
    assert "【下一步建议】" not in report
    from io import BytesIO
    from openpyxl import load_workbook
    workbook = load_workbook(BytesIO(artifact_svc.read_bytes(event["artifacts"][1]["id"])))
    sheet = workbook["测试清单"]
    assert [cell.value for cell in sheet[1]] == [
        "序号", "对象类型", "对象名称", "目标库", "校验结果", "原因分析", "修复建议", "转换SQL"
    ]
    rows_by_name = {sheet.cell(row=row, column=3).value: row for row in range(2, sheet.max_row + 1)}
    assert sheet.cell(row=rows_by_name["users"], column=2).value == "table"
    assert sheet.cell(row=rows_by_name["users"], column=5).value == "成功"
    assert sheet.cell(row=rows_by_name["user_orders"], column=2).value == "view"
    assert sheet.cell(row=rows_by_name["user_orders"], column=5).value == "失败"
    assert any(name and str(name).startswith("示例_") for name in rows_by_name)


def test_assessment_artifact_strips_output_file_section(tmp_path):
    """小评评估报告不携带输出文件清单，避免提示词格式泄漏进产物。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    reply = (
        "【分析结果】\n"
        "1. 兼容性分析\n"
        "- 数据库类型：Oracle 转 MySQL。\n"
        "- customers 表字段直接兼容。\n"
        "2. 风险点\n"
        "- Oracle 的 NUMBER 字段需确认 MySQL DECIMAL 精度。\n"
        "- 约束和索引需要在目标库重新验证。\n"
        "3. 输出文件\n"
        "- 迁移评估报告\n"
        "- 迁移脚本\n"
        "4. 下一步建议\n"
        "- 验证字段类型和长度映射。\n"
        "- 迁移后执行数据完整性校验。\n"
    )
    llm = ScriptedLLM([(reply, None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    custom_events = []
    for mode, payload in graph.stream(
            AgentState(user_input="执行迁移评估", conv_id=18,
                       agent_id="migration_commander",
                       task_state={"current_step": "ASSESSMENT_RUNNING",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom"]):
        if mode == "custom" and isinstance(payload, dict) and payload.get("type") == "sub_agent":
            custom_events.append(payload)

    event = custom_events[0]
    assert [f["path"].split("/")[-1] for f in event["artifacts"]] == ["评估报告.md"]
    report = artifact_svc.read(event["artifacts"][0]["id"])
    assert "Oracle 转 MySQL" in report
    assert "NUMBER 字段需确认 MySQL DECIMAL 精度" in report
    assert "验证字段类型和长度映射" in report
    assert "输出文件" not in report
    assert "迁移评估报告" not in report
    assert "迁移脚本" not in report


def test_capability_and_prerequisite_text_is_not_core_delivery():
    """职责说明、能力清单、材料前置条件不是智能体核心产物。"""
    from db_loc_agent.graph import _is_core_delivery

    samples = {
        "code_migrator": (
            "## 分析结果\n\n"
            "当前任务参数表明这是一次迁移评估任务，超出我（小改）的职责边界，本次不予执行。\n\n"
            "## 我能做什么\n\n"
            "- 驱动与连接改造：调整连接串与连接池配置\n\n"
            "## 需要你提供的材料\n\n"
            "1. 待改造的代码片段\n2. 源数据库类型与目标国产数据库类型"
        ),
        "migration_planner": (
            "## 编制落地方案前必须补充的输入\n\n"
            "以下评估与转换阶段产物未随任务传入，逐项列出：\n\n"
            "1. 实际源库类型及版本\n2. 兼容性评估结论\n\n"
            "## 请补充以下字段后，我将生成分步可执行的迁移落地实施方案"
        ),
        "migration_collector": (
            "## 采集服务说明\n\n"
            "我是小采，可以连接源库并采集表结构、索引、约束等元数据。\n\n"
            "## 需要你提供的材料\n\n"
            "- 源数据源连接信息\n- 对象范围说明"
        ),
    }
    for agent_id, text in samples.items():
        assert not _is_core_delivery(agent_id, text), (agent_id, text)


def test_stage_block_is_invisible_and_disables_next_agent_calls(tmp_path):
    """数据采集报告已生成时，执行节点完成；参数未齐仍停驻本阶段。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    llm = ScriptedLLM([("# 数据采集报告\n\n- 对象清单已采集", None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=4,
        agent_service=AgentService(db),
    )
    custom_events = []
    final = {}
    for mode, payload in graph.stream(
            AgentState(user_input="数据迁移", conv_id=12,
                       agent_id="migration_commander",
                       task_state={"current_step": "DATA_COLLECTION",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict):
            custom_events.append(payload)
        elif mode == "updates" and isinstance(payload, dict) and payload:
            update = next(iter(payload.values()), {})
            if isinstance(update, dict) and "agent_calls" in update:
                final = update

    assert [e.get("agent_id") for e in custom_events
            if e.get("event") == "calling_agent"] == ["migration_collector"]
    assert not any(e.get("event") == "stage_blocked" for e in custom_events)
    assert [c["agent_id"] for c in final["agent_calls"]] == ["migration_collector"]
    assert final["artifacts"]
    assert final["task_state"]["step_status"] == "done"
    assert final["task_state"]["awaiting_confirmation"] is False


def test_free_flow_allows_out_of_order_stage_agent(tmp_path):
    """关闭流程强校验后，指挥官可按用户指派提前调用后续阶段专家。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    llm = ScriptedLLM([("请先提供源库对象清单、目标库类型和停机窗口。", None)])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(
        user_input="不用先采集，直接评估", conv_id=14,
        agent_id="migration_assessor",
        task_state={"current_step": "DATA_COLLECTION",
                    "step_confirm": True, "pipeline_start": True,
                    "strict_flow": False}))

    assert [call["agent_id"] for call in final["agent_calls"]] == ["migration_assessor"]
    assert final["agent_calls"][0]["error"] == ""
    assert final["task_state"]["strict_flow"] is False


def test_commander_does_not_write_orchestration_index_artifact(tmp_path):
    """调度索引是过程数据，不是核心交付物，不得进入产物面板。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    call_tool = {"name": "call_agent",
                 "args": {"agent_id": "migration_assessor",
                          "question": "请评估迁移风险"},
                 "id": "call_no_index_1"}
    llm = ScriptedLLM([
        ("需要评估", [call_tool]),
        ("请先提供源库对象清单。", None),
        ("评估专家等待补充信息。", None),
    ])
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=3,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(
        user_input="继续评估", conv_id=11,
        agent_id="migration_commander",
        task_state={"current_step": "ASSESSMENT_RUNNING",
                    "step_confirm": True, "pipeline_start": True}))

    assert final.get("agent_calls")
    assert artifact_svc.list_namespaces(conv_id=11) == []


def test_direct_core_agent_writes_only_business_document(tmp_path):
    """用户直接咨询评估专家时，也只生成核心报告，不生成 result.json。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    report = "# 迁移评估报告\n\n兼容性结论：可以迁移，需处理函数差异。"

    class OneReplyLLM(FakeLLM):
        def parse(self, text):
            return {"intent": "migration_assess", "agent_id": "migration_assessor",
                    "object_name": "营业库", "source_db": "oracle", "target_db": "dm",
                    "skill_id": "", "confidence": 0.99}

    graph = build_graph(
        llm=OneReplyLLM(reply=report), rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=1,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(user_input="请评估营业库", conv_id=12))

    filenames = [f["path"].split("/")[-1] for f in final["artifacts"]]
    assert filenames == ["评估报告.md"]


def test_direct_collector_failure_report_creates_chinese_artifact(tmp_path):
    """用户直接指派小采时，结构化失败报告同样落为中文采集产物。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    report = (
        "# 数据采集报告\n\n"
        "## 采集状态：无法完成\n\n"
        "- 数据源 src_o 未配置\n"
        "- 对象清单采集数量：0"
    )

    class OneReplyLLM(FakeLLM):
        def parse(self, text):
            return {"intent": "migration_collect", "agent_id": "migration_collector",
                    "object_name": "src_o", "source_db": "oracle", "target_db": "dm",
                    "skill_id": "", "confidence": 0.99}

    graph = build_graph(
        llm=OneReplyLLM(reply=report), rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=1,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(user_input="请采集 src_o", conv_id=13))

    filenames = [f["path"].split("/")[-1] for f in final["artifacts"]]
    assert filenames == ["数据采集报告.md"]
    assert artifact_svc.read(final["artifacts"][0]["id"]) == report


def test_core_sql_artifacts_use_chinese_filenames(tmp_path):
    """新产物文件名使用中文业务名，技术对象名仅作为后缀保留。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    class ConverterLLM(FakeLLM):
        def parse(self, text):
            return {"intent": "ddl_convert", "agent_id": "ddl_converter",
                    "object_name": "USER_ORDER", "source_db": "oracle",
                    "target_db": "dm", "skill_id": "", "confidence": 0.99}

    graph = build_graph(
        llm=ConverterLLM(reply="# 转换报告\n\n```sql\nCREATE TABLE USER_ORDER (ID INT);\n```"),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=1,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(user_input="转换 USER_ORDER", conv_id=13))

    filenames = [f["path"].split("/")[-1] for f in final["artifacts"]]
    assert filenames == ["转换脚本_USER_ORDER.sql", "转换报告.md"]


def test_code_migrator_gets_role_boundary_and_ignores_mismatched_history(tmp_path):
    """用户指派小改时，不得因历史评估任务而代跑评估专家职责。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    prompt_svc = PromptService(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    class RecordingCodeLLM(FakeLLM):
        def __init__(self):
            super().__init__(reply="请切换到小评处理评估。")
            self.seen_messages = []

        def parse(self, text):
            return {"intent": "migration_assess", "agent_id": "migration_assessor",
                    "object_name": "", "source_db": "", "target_db": "",
                    "skill_id": "", "confidence": 0.95}

        def stream(self, messages, **kw):
            self.seen_messages.append(list(messages))
            yield from super().stream(messages, **kw)

    llm = RecordingCodeLLM()
    graph = build_graph(
        llm=llm, rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=prompt_svc, artifact_service=artifact_svc,
        skill_service=None, parse_schema={}, max_iterations=1,
        agent_service=AgentService(db),
    )
    final = graph.invoke(AgentState(
        user_input="干活", conv_id=14, agent_id="code_migrator",
        history=[{"role": "user", "content": "我要做数据迁移评估"},
                 {"role": "assistant", "content": "好的，我来帮你做数据迁移评估。"}]))

    assert final["agent_id"] == "code_migrator"
    system_prompt = str(llm.seen_messages[0][0].content)
    assert "职责边界（最高优先级）" in system_prompt
    assert "不得执行迁移评估" in system_prompt
    assert "小评" in system_prompt
    assert "不要把历史会话中的迁移评估任务当作当前任务" in system_prompt


def test_graph_exposes_six_agent_nodes(tmp_path):
    """主图节点是六个迁移智能体，不再是 parse/route/execute 壳层。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)

    graph = build_graph(
        llm=FakeLLM(),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=PromptService(db),
        artifact_service=ArtifactService(db, str(tmp_path / "artifacts")),
        skill_service=None,
        parse_schema={},
        agent_service=AgentService(db),
    )
    nodes = set(graph.get_graph().nodes)
    migration_nodes = {
        "migration_commander", "migration_collector", "migration_assessor",
        "migration_planner", "code_migrator", "migration_verifier",
    }
    assert migration_nodes.issubset(nodes)
    assert not (nodes - migration_nodes - {"__start__", "__end__"})
    assert not {"parse", "route", "execute", "artifact", "reply",
                "specialist_agent"} & nodes


def test_route_agent_maps_current_stage_and_finished_state():
    rt = _runtime_for_routing()
    cases = {
        "DATA_COLLECTION": "migration_collector",
        "ASSESSMENT_RUNNING": "migration_assessor",
        "SOLUTION_GENERATE": "migration_planner",
        "CODE_MIGRATE": "code_migrator",
        "VERIFY_STAGE": "migration_verifier",
    }
    for stage, agent_id in cases.items():
        state = AgentState(agent_id="migration_commander",
                           task_state={"current_step": stage})
        assert rt.route_agent(state) == agent_id
    assert rt.route_agent(AgentState(
        agent_id="migration_commander",
        task_state={"current_step": "FINISHED"})) == "__end__"


def test_route_agent_enforces_explicit_stage_assignment():
    rt = _runtime_for_routing()
    blocked = AgentState(
        agent_id="migration_assessor",
        task_state={"current_step": "DATA_COLLECTION", "strict_flow": True})
    assert rt.route_agent(blocked) == "migration_commander"
    assert "当前阶段" in blocked["route_reason"]

    free = AgentState(
        agent_id="migration_assessor",
        task_state={"current_step": "DATA_COLLECTION", "strict_flow": False})
    assert rt.route_agent(free) == "migration_assessor"


def test_should_advance_distinguishes_stay_wait_next_and_back():
    rt = _runtime_for_routing()
    complete_params = {
        "source_ds_id": "src_o", "source_db_type": "oracle",
        "target_db_type": "dm", "db_business_level": "核心",
        "db_capacity": "500GB", "stop_window": "4小时",
    }
    assert rt.should_advance(AgentState(task_state={
        "current_step": "DATA_COLLECTION", "task_params": {}})) == "stay"
    assert rt.should_advance(AgentState(task_state={
        "current_step": "DATA_COLLECTION",
        "task_params": complete_params})) == "wait_user"
    assert rt.should_advance(AgentState(task_state={
        "current_step": "DATA_COLLECTION",
        "task_params": {**complete_params, "task_confirm": "确认"}})) == "next"
    assert rt.should_advance(AgentState(task_state={
        "current_step": "DATA_COLLECTION", "step_confirm": False,
        "task_params": complete_params})) == "next"
    assert rt.should_advance(AgentState(task_state={
        "current_step": "VERIFY_STAGE", "verification_failed": True,
        "task_params": {"task_confirm": "确认"}})) == "back"


def test_agent_graph_stage_node_creates_artifact_and_waits(tmp_path):
    """阶段节点执行小评；核心报告入产物，图停驻等待用户确认。"""
    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))
    report = "# 迁移评估报告\n\n兼容性风险为中，主要集中在函数兼容性。"

    class CollectorStageLLM(FakeLLM):
        def parse(self, text):
            return {"intent": "migration_overall", "agent_id": "migration_commander",
                    "object_name": "营业库", "source_db": "oracle", "target_db": "dm",
                    "skill_id": "", "confidence": 1.0}

    graph = build_graph(
        llm=CollectorStageLLM(reply=report),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=PromptService(db),
        artifact_service=artifact_svc,
        skill_service=None,
        parse_schema={},
        agent_service=AgentService(db),
    )
    custom_events = []
    final = {}
    for mode, payload in graph.stream(
            AgentState(user_input="继续评估", conv_id=20,
                       agent_id="migration_commander",
                       task_state={"current_step": "ASSESSMENT_RUNNING",
                                   "step_confirm": True, "pipeline_start": True}),
            stream_mode=["custom", "updates"]):
        if mode == "custom" and isinstance(payload, dict):
            custom_events.append(payload)
        elif mode == "updates":
            final.update(next(iter(payload.values()), {}))

    sub_events = [e for e in custom_events if e.get("type") == "sub_agent"]
    assert [e["agent_id"] for e in sub_events] == ["migration_assessor"]
    assert [f["path"].split("/")[-1] for f in sub_events[0]["artifacts"]] == ["评估报告.md"]
    assert final["task_state"]["step_status"] == "waiting_confirm"
    assert final["task_state"]["awaiting_confirmation"] is True


def test_excel_collection_runs_deterministically_without_llm(tmp_path):
    """Excel 采集必须本地扫描工作表，不能等待模型自觉调用工具。"""
    from openpyxl import Workbook

    from db_loc_agent.db_tools import DbTools

    db = Database(str(tmp_path / "t.db"))
    from db_loc_agent.bootstrap import bootstrap
    bootstrap(db)
    artifact_svc = ArtifactService(db, str(tmp_path / "artifacts"))

    workbook = Workbook()
    workbook.active.title = "用户表"
    workbook.active.append(["id", "名称"])
    workbook.active.append([1, "移动"])
    workbook.create_sheet("订单表").append(["id", "金额"])
    excel_path = tmp_path / "采集.xlsx"
    workbook.save(excel_path)

    db_tools = DbTools(db)
    db_tools.create("src_excel", kind="excel", location=str(excel_path))

    class NoInvokeLLM(FakeLLM):
        def invoke(self, messages, **kwargs):
            raise AssertionError("Excel 数据采集不应依赖 LLM invoke")

        def stream(self, messages, **kwargs):
            raise AssertionError("Excel 数据采集不应依赖 LLM stream")

        def parse(self, text):
            return {"intent": "migration_collect", "agent_id": "migration_collector",
                    "object_name": "营业库", "source_db": "excel", "target_db": "mysql",
                    "skill_id": "", "confidence": 1.0}

    graph = build_graph(
        llm=NoInvokeLLM(),
        rag_search=lambda q, top_k=4, space_id=None: [],
        prompt_service=PromptService(db),
        artifact_service=artifact_svc,
        skill_service=None,
        parse_schema={},
        agent_service=AgentService(db),
        db_tools=db_tools,
    )

    final = graph.invoke(AgentState(
        user_input="开始采集", conv_id=21,
        agent_id="migration_commander",
        task_state={
            "current_step": "DATA_COLLECTION",
            "pipeline_start": True,
            "step_confirm": True,
            "task_params": {
                "source_ds_id": "src_excel",
                "source_db_type": "excel",
                "target_db_type": "mysql",
                "db_business_level": "核心",
                "db_capacity": "500GB",
            },
        },
    ))

    namespaces = artifact_svc.list_namespaces(conv_id=21)
    assert len(namespaces) == 1
    artifacts = artifact_svc.list_files(namespaces[0])
    assert len(artifacts) == 1
    report_path = Path(artifacts[0]["path"])
    report = report_path.read_text(encoding="utf-8")
    assert "用户表" in report
    assert "订单表" in report
    assert "示例补充对象（演示数据）" in report
    assert "V_ORDER_SUMMARY" in report
    assert "P_CALC_USER_BALANCE" in report
    assert "示例对象未计入真实对象数" in report
    assert "待确认信息" not in report
    assert final["task_state"]["task_params"]["object_count"] == 2
    assert final["task_state"]["step_status"] == "done"
